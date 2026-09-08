import sys
import os

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

doc = Document()

# Title
title = doc.add_heading('Comprehensive ERP Architecture: Marg ERP vs. Enterprise Systems', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('This document provides an advanced architectural comparison between high-speed SME systems (Marg ERP) and Top-Tier Enterprise ERPs (SAP S/4HANA), utilizing workflow diagrams to map the exact lifecycle of Challans, Invoices, and Billing Series.')

# ---------------------------------------------------------
# Section 1: The Marg ERP Workflow (Challans & Series)
# ---------------------------------------------------------
doc.add_heading('1. The Marg ERP Workflow: Challans to Bills & Series Logic', level=1)
doc.add_paragraph('Marg ERP dominates the Pharma and FMCG distribution sectors due to its extremely fast, keyboard-driven workflows. Two of its most defining features are the "Challan Conversion" and "Billing Series".')

doc.add_heading('A. The Marg Challan System (Delivery Note)', level=2)
doc.add_paragraph('In Marg, businesses often ship goods before finalizing the tax invoice. They generate a "Delivery Challan".')
doc.add_paragraph('  • Stock Impact: The moment a Challan is saved, the physical inventory drops.')
doc.add_paragraph('  • Financial Impact: It does NOT hit the final Accounts Receivable or Revenue ledgers.')
doc.add_paragraph('  • Conversion: Later, users open a "New Sale Bill", press a hotkey to "Load Challans", and select multiple pending challans to merge into a single final Tax Invoice.')

doc.add_heading('Workflow Diagram: Marg Challan to Bill', level=2)
p_diag1 = doc.add_paragraph()
p_diag1.add_run("""
[ Sales Order ]
       │
       ▼
[ Delivery Challan 1 ] ─────┐ (Inventory drops by 50)
                            │
[ Delivery Challan 2 ] ─────┼───> [ Convert to Sale Bill ] 
                            │     (Consolidates to 150 units)
[ Delivery Challan 3 ] ─────┘     (Hits AR, Revenue & GST Ledgers)
""")

doc.add_heading('B. Marg Billing Series', level=2)
doc.add_paragraph('Marg organizes invoices using "Series" (e.g., Series A, Series B). This allows a single company to maintain parallel invoice sequences.')
doc.add_paragraph('  • Series A: Used for local GST B2B billing (Sequence: A-001, A-002).')
doc.add_paragraph('  • Series B: Used for Cash Retail billing (Sequence: B-001, B-002).')

# ---------------------------------------------------------
# Section 2: How Enterprise ERPs Handle This (SAP/Oracle)
# ---------------------------------------------------------
doc.add_heading('2. Enterprise Equivalents: SAP & Oracle NetSuite', level=1)
doc.add_paragraph('While Marg relies on the "Challan" as a standalone entry screen, Enterprise ERPs build this directly into the core logistical architecture.')

doc.add_heading('A. The Enterprise "Challan" = Outbound Delivery & PGI', level=2)
doc.add_paragraph('In SAP, the equivalent of a Marg Challan is the "Outbound Delivery Document". However, enterprise systems enforce strict dual-entry accounting even at the delivery stage.')
doc.add_paragraph('  • Marg Challan simply drops stock. SAP Outbound Delivery with Post Goods Issue (PGI) triggers a massive accounting event: Debit COGS, Credit Inventory Asset.')
doc.add_paragraph('  • Enterprise systems allow "Collective Billing" where multiple Delivery Documents are automatically rolled up into a single Billing Document overnight.')

doc.add_heading('B. The Enterprise "Series" = Document Types & Number Ranges', level=2)
doc.add_paragraph('SAP does not use hardcoded "Series". Instead, it uses a highly scalable matrix:')
doc.add_paragraph('  • Document Type (e.g., RV for Standard Invoice, RE for Return).')
doc.add_paragraph('  • Number Range Object (Assigns a sequence based on the Company Code and Year, e.g., 2025-RV-0001).')

# ---------------------------------------------------------
# Section 3: Visual Workflow & GL Impact Comparison
# ---------------------------------------------------------
doc.add_heading('3. Side-by-Side Workflow & Ledger Diagrams', level=1)

table = doc.add_table(rows=1, cols=2)
table.style = 'Table Grid'
hdr_cells = table.rows[0].cells
hdr_cells[0].text = 'Marg ERP Workflow (Challan System)'
hdr_cells[1].text = 'Enterprise ERP Workflow (PGI System)'

row_cells = table.add_row().cells
row_cells[0].text = "STEP 1: ISSUE CHALLAN\nUser creates a Delivery Challan.\n\nInventory: Drops.\nFinance: No GL entry posted."
row_cells[1].text = "STEP 1: POST GOODS ISSUE (PGI)\nWarehouse ships the Outbound Delivery.\n\nInventory: Drops.\nFinance: Auto-posts Journal.\n  [Dr] Cost of Goods Sold\n  [Cr] Inventory Asset"

row_cells = table.add_row().cells
row_cells[0].text = "STEP 2: CONVERT TO BILL\nUser loads Challan into a Sale Bill.\n\nFinance: Auto-posts Journal.\n  [Dr] Customer AR\n  [Cr] Sales Revenue\n  [Cr] Output Tax"
row_cells[1].text = "STEP 2: BILLING DOCUMENT\nSystem auto-generates Invoice from Delivery.\n\nFinance: Auto-posts Journal.\n  [Dr] Customer AR\n  [Cr] Sales Revenue\n  [Cr] Output Tax"

# ---------------------------------------------------------
# Section 4: Advanced Procure-to-Pay (P2P) Diagram
# ---------------------------------------------------------
doc.add_heading('4. Procure-to-Pay (P2P) Visual Workflow', level=1)
doc.add_paragraph('This diagram illustrates the 3-Way Match suspense accounting workflow used to prevent vendor fraud.')

p_diag2 = doc.add_paragraph()
p_diag2.add_run("""
┌─────────────────────────┐
│ 1. PURCHASE ORDER (PO)  │ -> Commitment generated. No GL impact.
└───────────┬─────────────┘
            │
            ▼
┌─────────────────────────┐    [Dr] Inventory Asset
│ 2. GOODS RECEIPT (Dock) │ -> [Cr] GR/IR Suspense Account
└───────────┬─────────────┘    (Stock arrives, but no bill yet)
            │
            ▼
┌─────────────────────────┐    [Dr] GR/IR Suspense Account (Zeros out)
│ 3. 3-WAY MATCH (MIRO)   │ -> [Dr] Input Tax (GST)
└───────────┬─────────────┘    [Cr] Vendor Accounts Payable
            │
            ▼
┌─────────────────────────┐    [Dr] Vendor Accounts Payable
│ 4. VENDOR PAYMENT       │ -> [Cr] Bank Account
└─────────────────────────┘
""")

folder_path = r'C:\Users\DELL\OneDrive\Desktop\erp2\billing'
file_path = os.path.join(folder_path, 'Enterprise_Sales_Purchase_Billing_Workflow_v2.docx')
doc.save(file_path)
print('Document successfully rewritten with Diagrams, Marg Comparisons, and Workflows (v2)!')
