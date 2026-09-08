import sys
import os

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Enterprise Finance, Sales & Procurement Master Architecture', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('A comprehensive, unified blueprint combining top-tier Enterprise ERP architectures (SAP/Oracle) with essential high-speed SME workflows (Marg Challans, Tally Outstanding tracking). This dictates the exact hierarchies, workflows, and ledger impacts for the system.')

# ---------------------------------------------------------
# 1. Architectural Hierarchy
# ---------------------------------------------------------
doc.add_heading('1. Core Architectural Hierarchy', level=1)
doc.add_paragraph('To maintain scale and audit compliance, the system enforces a strict Separation of Concerns across three layers:')
doc.add_paragraph('  1. Commercial Intent: Orders & Quotes. These reserve stock and budget but do not post to the General Ledger (GL).')
doc.add_paragraph('  2. Physical Logistics: Deliveries & Receipts. These alter physical inventory and hit Inventory Valuation/COGS, but defer Revenue and Vendor Liability.')
doc.add_paragraph('  3. Financial Accounting: Invoices & Billing. These establish the legal liability, hitting Accounts Receivable/Payable, Revenue, and Tax.')

# ---------------------------------------------------------
# 2. Order-to-Cash (O2C) & Sales Workflow
# ---------------------------------------------------------
doc.add_heading('2. Order-to-Cash (Sales & Logistics Workflow)', level=1)
doc.add_paragraph('The system supports both standard Enterprise flows and flexible SME Challan flows.')

doc.add_heading('A. The Standard Enterprise Flow (PGI to Billing)', level=2)
doc.add_paragraph('Step 1: Sales Order -> Stock is soft-allocated by the ATP engine. [No GL Impact]')
doc.add_paragraph('Step 2: Post Goods Issue (PGI) -> The box is shipped.')
doc.add_paragraph('   • GL Impact: [Debit] Cost of Goods Sold | [Credit] Inventory Asset')
doc.add_paragraph('Step 3: Billing Document -> The final tax invoice is generated.')
doc.add_paragraph('   • GL Impact: [Debit] Customer AR | [Credit] Sales Revenue | [Credit] Output Tax')

doc.add_heading('B. The SME Flexible Flow: Marg-Style Challans & Series (Targeted Feature)', level=2)
doc.add_paragraph('To accommodate FMCG/Pharma workflows where goods ship before final pricing, we incorporate Challan conversion and flexible Series:')
doc.add_paragraph('  • Delivery Challan: Ships the goods and drops inventory. Defers AR/Revenue posting.')
doc.add_paragraph('  • Load Challans to Bill: Users can open a "Sale Bill", select multiple pending Delivery Challans, and merge them into a single consolidated Tax Invoice.')
doc.add_paragraph('  • Billing Series: Users can define custom prefixes (e.g., Series A for Wholesale, Series B for Retail) to run parallel, legally compliant invoice sequences.')

# ---------------------------------------------------------
# 3. Procure-to-Pay (P2P) & Purchasing Workflow
# ---------------------------------------------------------
doc.add_heading('3. Procure-to-Pay (Purchasing Workflow)', level=1)
doc.add_paragraph('Utilizing SAP-style Suspense Accounting to handle timing mismatches between the warehouse and the accounts payable desk.')

doc.add_paragraph('Step 1: Purchase Order -> Encumbers budget. [No GL Impact]')
doc.add_paragraph('Step 2: Goods Receipt (GR) -> The truck arrives at the dock, but no bill is provided yet.')
doc.add_paragraph('   • GL Impact: [Debit] Inventory Asset | [Credit] GR/IR Suspense Account')
doc.add_paragraph('   • Note: The GR/IR (Goods Receipt/Invoice Receipt) clearing account prevents us from crediting the vendor prematurely.')
doc.add_paragraph('Step 3: 3-Way Match Invoice (MIRO) -> The vendor emails the bill. The system verifies PO = GR = Invoice.')
doc.add_paragraph('   • GL Impact: [Debit] GR/IR Suspense (Zeros out) | [Debit] Input Tax | [Credit] Vendor AP')

# ---------------------------------------------------------
# 4. Outstanding Tracking: Tally-Style Bill-by-Bill
# ---------------------------------------------------------
doc.add_heading('4. Outstanding Management (Bill-by-Bill Adjustments)', level=1)
doc.add_paragraph('Standard systems track lump-sum balances. Our system tracks exact invoice matching:')
doc.add_paragraph('  • Against Reference: When a customer pays , the Receipt Voucher forces the user to apply it against a specific invoice (e.g., INV-1045).')
doc.add_paragraph('  • Architecture: An internal invoice_allocations bridge table links the Receipt ID to the Sales Invoice ID, enabling perfect 30/60/90-Day Aging reports.')

# ---------------------------------------------------------
# 5. Dynamic Billing & Pricing Engine
# ---------------------------------------------------------
doc.add_heading('5. The Dynamic Pricing Engine (Condition Technique)', level=1)
doc.add_paragraph('Pricing is decoupled from the product master, preventing hard-coded accounting errors.')
doc.add_paragraph('  • Access Sequence: Automatically searches for Customer-Specific discounts before falling back to standard rates.')
doc.add_paragraph('  • Account Keys: The engine tags lines with keys (e.g., ERL for Revenue, MWS for Tax). At the moment of saving, an Account Determination Matrix routes the ERL value to GL 400000 and MWS to GL 220000, allowing Marketing to change discounts without breaking Finance.')

# ---------------------------------------------------------
# 6. Transaction Integrity (DAG & Immutability)
# ---------------------------------------------------------
doc.add_heading('6. Transaction Integrity & Audit Logging', level=1)
doc.add_paragraph('  • Directed Acyclic Graph (DAG): Documents are permanently linked. A quote links to an order, which links to a delivery, which links to a bill.')
doc.add_paragraph('  • Immutability & Storno: Hard deletes are banned. If a bill is wrong, the system generates a Reversal Journal (Negative Debits/Credits) to maintain a flawless audit trail.')

folder_path = r'C:\Users\DELL\OneDrive\Desktop\erp2\billing'
file_path = os.path.join(folder_path, 'Enterprise_Architecture_Master_Blueprint_v3.docx')
doc.save(file_path)
print('Master Blueprint v3 generated successfully!')
