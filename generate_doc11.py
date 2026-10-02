import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc11():
    doc = docx.Document()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DOC-11\nPRODUCT, SKU & BATCH MASTER FRAMEWORK\n")
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.font.bold = True
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = subtitle.add_run("Implementation Specification • Phase 4: Master Data")
    s_run.font.name = 'Arial'
    s_run.font.size = Pt(12)
    s_run.font.italic = True
    
    doc.add_page_break()
    
    sections = [
        ("1. Executive Summary", 
         "The Product and SKU Master is the central nervous system of inventory and billing. Every purchase bill, sales invoice, stock adjustment, and profitability report relies on the structural integrity of this master data. The ERP must model items not just as text strings, but as multi-dimensional financial and logistical assets. Because this is a C&F-native ERP, the system must strictly enforce Batch, Expiry, and MRP tracking at the structural level."),
         
        ("2. Relationship to the ERP Company Architecture",
         "The Item Master is company-scoped. However, for C&F agents operating under a Principal Company, the Platform Admin can optionally push a 'Global Principal Catalog' down to the client company to prevent manual data entry errors. The client company can then augment this catalog with local warehouse rules and pricing overrides.\n\nPLATFORM (Principal Catalog)\n  ↓\nCLIENT COMPANY (Local Overrides & Tax)\n  ↓\nITEM / SKU\n  ├─ Base Identity (Name, HSN, Tax)\n  ├─ Logistics (UoM, Weight, CBM)\n  ├─ Traceability (Batch/Expiry Rules)\n  └─ Financials (MRP, Pricing, Accounts)"),
         
        ("3. Item Master Philosophy",
         "3.1 Base Product vs SKU Variant\nUnlike simple retail software, a robust ERP separates the 'Base Product' from the 'SKU Variant'. However, for C&F operations which typically deal in finished goods, the system will optimize for speed by flattening this into a unified Item Master, while relying on Batch/Lot tracking to differentiate specific manufacturing runs.\n\n3.2 Zero Duplicate Entry\nIf an item is purchased, the Purchase Bill pulls its HSN, Tax Rate, and Default UoM from the master. The user must never manually type tax rates during transactions. If statutory rates change, they are updated in the master with an 'Effective Date', ensuring historical invoices remain mathematically immutable.\n\n3.3 Traceability is Non-Negotiable\nFor pharmaceutical, FMCG, and C&F sectors, traceability is legally required. The item master defines *how* an item is tracked (Batch No, Serial No, Expiry Date, or None). The transaction engine (DOC-19) strictly enforces these rules."),
         
        ("4. Product Identity & Hierarchy",
         "4.1 Item Categorization\nItems must be grouped for MIS reporting and ledger mapping. The hierarchy is:\nPRINCIPAL (Manufacturer) → ITEM GROUP (Category) → SUB-GROUP → ITEM.\n\n4.2 Naming Conventions & Search\nThe system must maintain a 'Strict Legal Name' (for tax invoices) and a 'Trade Name / Short Name' (for fast keyboard searching). Duplicate prevention operates asynchronously during creation, alerting the user if an item with a similar name or Principal Barcode already exists."),
         
        ("5. Product Master — Complete Functional Structure",
         "5.1 Product Master Screen UI Mockup\n\n┌──────────────────────────────────────────────────────────────────────┐\n│ ITEM MASTER  [Active ▼]                 Code: ITM-88942             │\n├──────────────────────────────────────────────────────────────────────┤\n│ LEGAL NAME: Paracetamol 500mg (10x10) Strip                         │\n│ PRINCIPAL: PharmaCorp Ltd.         GROUP: Analgesics                │\n├──────────────────────────────────────────────────────────────────────┤\n│ GENERAL | UoM & LOGISTICS | TAX & HSN | BATCH RULES | PRICING       │\n├──────────────────────────────────────────────────────────────────────┤\n│ Base UoM: STRIP     Weight: 0.05 KG     Tracking: [Batch + Expiry]  │\n├──────────────────────────────────────────────────────────────────────┤\n│ Default Sale A/C: Sales - 12%      Default Purch A/C: Purch - 12%   │\n├──────────────────────────────────────────────────────────────────────┤\n│ Stock Available: 4,500 STRIPS      Reorder Level: 500 STRIPS        │\n└──────────────────────────────────────────────────────────────────────┘\n"),
         
        ("6. Unit of Measure (UoM) & Logistics",
         "6.1 Multi-UoM Architecture\nC&F billing frequently purchases in BOXES and sells in PIECES or STRIPS. The item master must define the Base UoM (the smallest tracked unit) and Alternate UoMs.\n\nExample:\n- Base UoM: PIECE (Factor 1)\n- Alternate 1: BOX (Factor 10 Pieces)\n- Alternate 2: CASE (Factor 100 Pieces)\n\nAll stock ledger balances are stored in the Base UoM. Transactions can be entered in Alternate UoMs, and the ERP will mathematically resolve them to the Base UoM seamlessly.\n\n6.2 Weight & Volume (CBM)\nFor the Cargo & Consignment module (DOC-23), every item must have a defined Gross Weight and Volume. When a sales invoice is generated, the ERP will automatically calculate the total Consignment Weight based on these master data fields, driving the Freight Cost engine."),
         
        ("7. Tax, HSN & Statutory Details",
         "7.1 HSN / SAC Enforcement\nEvery item must have a valid HSN (Harmonized System of Nomenclature) code. The system will maintain a centralized master list of HSN codes. The Item Master links to this list rather than storing free-text.\n\n7.2 Tax Rate Derivation\nTax rates (GST 5%, 12%, 18%, 28%) are tied to the HSN code or explicitly set on the Item Master. The billing engine will pull this rate automatically. If an item is tax-exempt, it must be explicitly flagged with an Exemption Reason Code for statutory reporting."),
         
        ("8. Batch, Expiry & Serial Number Architecture",
         "This is a massive differentiator for this ERP. The Item Master does not store the batches; it stores the *Rules* for the batches.\n\n8.1 Tracking Modes\n- None: Standard items (e.g., promotional materials).\n- Batch Only: Items tracked by a manufacturer lot number.\n- Batch + Expiry: FMCG/Pharma items. The system will enforce FEFO (First Expiry First Out) picking logic during Dispatch (DOC-25).\n- Serialized: Electronics/High-value items tracked uniquely per piece.\n\n8.2 Expiry Alerts\nIf an item is marked as Expiry Tracked, the master defines the 'Shelf Life Days' and 'Alert Before Expiry Days'. The notification engine will generate alerts to the warehouse manager when stock enters the 'Near Expiry' threshold."),
         
        ("9. Pricing & Smart MRP Strategy",
         "9.1 Price Lists\nAn item does not have a single 'Price'. It has Price Lists. \n- Standard Selling Price\n- Dealer Price\n- Institutional Price\n- Maximum Retail Price (MRP)\n\n9.2 Smart MRP (C&F Specific)\nIn FMCG/Pharma, the same product can have different MRPs based on the manufacturing batch. The Item Master defines the 'Base Prices', but the active Selling Price is dynamically resolved during billing based on the selected Batch Number and Customer Price List. The ERP will actively block billing if a user attempts to sell an item above its active Batch MRP."),
         
        ("10. Financial Linkage (General Ledger)",
         "Inheriting from DOC-08 (Posting Engine) and DOC-16 (Chart of Accounts), the Item Master must dictate its financial footprint.\n- Default Sales Account (e.g., Sales @ 12%)\n- Default Purchase Account (e.g., Purchases @ 12%)\n- Default Inventory Account (e.g., Stock-in-Trade)\n- Default COGS Account\nUsers must never manually select these ledgers during a fast-paced billing operation; the system handles the double-entry automatically based on the Item Master mapping."),
         
        ("11. Warehouse & Storage Defaults",
         "For multi-warehouse operations (DOC-12), the Item Master can define a Default Receiving Warehouse and a Default Bin/Rack Location. This optimizes the Inward/GRN process by telling the warehouse staff exactly where to place incoming stock."),
         
        ("12. Keyboard-First UX Workflows",
         "Following DOC-05, creating an Item Master must be lightning-fast. The UI will use progressive disclosure. The user inputs Name, selects Principal, hits Tab, inputs HSN, hits Enter to save defaults. Advanced configurations (Multi-UoM, Accounting Overrides) are tucked away in secondary tabs accessible via keyboard shortcuts (e.g., Alt+2)."),
         
        ("13. Active/Inactive & Archiving Lifecycle",
         "Items cannot be hard-deleted once stock has been initialized or transactions exist. They are marked 'Inactive'. Inactive items are hidden from standard Sales/Purchase lookups but remain in historical reports and the General Ledger audit trail."),
         
        ("14. Local-First Synchronization Rules",
         "Master data changes made on a LAN node are synchronized to the Cloud Light Database asynchronously. Master Sync uses Last-Write-Wins (LWW). Since Item Masters have a low collision probability compared to transactions, eventual consistency is perfectly safe."),
         
        ("15. Failure & Edge Cases",
         "15.1 Changing UoM post-transaction\nIf an item has historical stock movements, the Base UoM is strictly locked. Altering it would corrupt historical valuations. If a fundamental change is required, the user must create a new Item Master and transfer stock via a controlled adjustment.\n\n15.2 Missing HSN during Billing\nIf an item somehow bypasses validation and lacks an HSN code, the Billing Engine (DOC-22) will physically block the generation of a Tax Invoice until the Master is corrected, ensuring compliance."),
         
        ("16. Implementation Checklist",
         "[ ] Database schema migrations (items, item_uoms, item_prices)\n[ ] Unified API routes with Pydantic validation\n[ ] Multi-UoM conversion middleware\n[ ] HSN/Tax integration layer\n[ ] Batch tracking rule enforcement module\n[ ] Keyboard-first React UI")
    ]
    
    for heading, text in sections:
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)

    # Save
    out_path = r'C:\Users\DELL\Desktop\New folder (4)\DOC-11_Product_and_SKU_Master_v1_0_RELEASE.docx'
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_doc11()
