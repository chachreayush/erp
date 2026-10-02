from docx import Document

def create_docx():
    doc = Document()
    doc.add_heading('Software Requirements & Architecture Specification', 0)

    doc.add_heading('1. Executive Summary', level=1)
    doc.add_paragraph('This document outlines the architectural requirements and comprehensive roadmap for the ERP System. The software is designed as a highly scalable, multi-tenant application that prioritizes data sovereignty and speed through a local-first architecture. It operates independently on local hardware while maintaining a Live Database Sync globally using cloud storage infrastructure, creating a seamless bridge between local networks and remote clients.')

    doc.add_heading('2. Core Architectural Pillars', level=1)
    
    doc.add_heading('2.1 Multi-Tenant Architecture', level=2)
    doc.add_paragraph('The system natively supports multiple organizations and tenants within a single deployment. Every record in the database is strictly tied to an organization ID. Role-Based Access Control (RBAC) ensures data is isolated and secure.')

    doc.add_heading('2.2 Local-First Execution & LAN Clients (Heavy Transactions)', level=2)
    doc.add_paragraph('The primary node of the software runs directly on the client\'s local machine or local server. Within the primary office, multiple clients can connect to the local server over a Local Area Network (LAN). These LAN clients are responsible for heavy data entry and high-volume transactions. The local server acts as the single source of truth, employing advanced concurrency handling to prevent duplicate data entries and collisions when multiple users submit data simultaneously.')

    doc.add_heading('2.3 Live Database Sync & Cloud Access', level=2)
    doc.add_paragraph('To support remote work and external clients, the system employs a real-time, live synchronization strategy rather than a batch process. Similar to a live continuous deployment pipeline, the local database actively pushes state changes in real-time to the "light database" hosted on cloud storage. Remote users access the software using this live cloud state to perform lighter tasks. Any light data entry performed remotely is pushed to the cloud state and instantly mirrored back to the local primary server, ensuring absolute consistency without data duplication.')

    doc.add_heading('3. Comprehensive Feature Modules (Built & Planned)', level=1)
    
    doc.add_heading('3.1 Financial & Accounting Engine', level=2)
    doc.add_paragraph('- Indian Taxation Account Head Hierarchy: A strict 3-level hierarchy mapping to Assets, Liabilities, Income, and Expenses.\n'
                      '- System-Protected Ledgers: Core ledgers are permanently protected from deletion to preserve accounting integrity.\n'
                      '- Bill-by-Bill Allocation: Precise mapping of receipts to outstanding invoices.')
                      
    doc.add_heading('3.2 Operational Workflows', level=2)
    doc.add_paragraph('- Challan-to-Invoice Data Pipeline: A seamless conversion pipeline that maps Delivery Challans directly into Tax Invoices.\n'
                      '- Smart MRP: Intelligent algorithms to suggest and apply optimal retail prices.\n'
                      '- Crash Recovery Engine: The system constantly saves in-progress forms as drafts. Users can resume exactly where they left off after a crash.')

    doc.add_heading('3.3 CRM & Field Management', level=2)
    doc.add_paragraph('- Field Staff Tracking & Scheduling: A comprehensive module designed to track field staff operations on a day-to-day, month-to-month, and year-to-year basis.\n'
                      '- Bulletin Board: A unified messaging center for internal announcements.\n'
                      '- Permissions Matrix: Detailed grids allowing administrators to toggle granular permissions.')

    doc.add_heading('4. Technology Stack', level=1)
    doc.add_paragraph('Frontend: React, TypeScript, Vite, TailwindCSS (Dark Enterprise Theme)\n'
                      'Desktop Runtime: Tauri (Rust-based local execution)\n'
                      'Backend API: FastAPI (Python)\n'
                      'Database: SQLAlchemy ORM (SQLite / PostgreSQL for LAN environments)')

    doc.add_heading('5. Implementation Roadmap', level=1)
    doc.add_paragraph('1. Phase 1: Core ERP, Master Ledgers, Inventory, and Multi-tenancy (Completed)\n'
                      '2. Phase 2: Financial Accounting Engine and Indian Taxation Structure (Completed)\n'
                      '3. Phase 3: Operational Workflows (Challan-to-Invoice Pipeline, CRM) (In Progress)\n'
                      '4. Phase 4: Field Staff Tracking & Scheduling Module (Upcoming)\n'
                      '5. Phase 5: Live Cloud Storage Sync Engine & LAN Optimization (Upcoming)')

    doc.save(r'c:\Users\DELL\OneDrive\Desktop\erp2\Architecture_Requirements.docx')
    print("Document updated successfully at c:\\Users\\DELL\\OneDrive\\Desktop\\erp2\\Architecture_Requirements.docx")

if __name__ == "__main__":
    create_docx()
