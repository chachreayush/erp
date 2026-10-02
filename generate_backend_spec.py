import docx
from docx.shared import Pt, RGBColor

doc = docx.Document()

# Title
doc.add_heading('BACKEND API & DATABASE IMPLEMENTATION SPECIFICATION (For Claude)', 0)

# Section 1: Bill-by-Bill Allocation Engine
doc.add_heading('1. Bill-by-Bill Allocation Engine (Receipts & Payments)', level=1)
doc.add_paragraph('This section dictates the strict backend implementation for the Bill-by-Bill allocation engine, replicating the Marg/Tally logic for mapping generic Receipt/Payment Vouchers to specific transactions.')

doc.add_heading('1.1 Core Business Rules & Floating Vouchers', level=2)
doc.add_paragraph('1. When a user enters a Receipt Voucher, the API MUST allow them to allocate the amount against multiple target types: Sales Invoices, Credit Notes (CN), Debit Notes (DN), and previously unallocated Floating Vouchers.')
doc.add_paragraph('2. Floating Vouchers: If the user does not allocate the entire receipt amount, the remaining balance MUST be stored as a "Floating" or "On Account" advance. This floating balance can be pulled up and allocated against future invoices. If the user explicitly leaves it floating, the backend must save it with `is_floating=True`.')
doc.add_paragraph('3. Append-Only/Offset Rule: Once an allocation is saved, it cannot be hard-deleted. Mistakes must be corrected by creating negative offset allocations.')

doc.add_heading('1.2 Database Schema (SQLAlchemy)', level=2)
code = doc.add_paragraph()
code.add_run("""class InvoiceAllocation(Base):
    __tablename__ = 'invoice_allocations'
    id = Column(UUID(as_uuid=True), primary_key=True, default=uuid.uuid4)
    organization_id = Column(UUID(as_uuid=True), index=True)
    source_voucher_id = Column(UUID(as_uuid=True)) # The Receipt or Payment Voucher
    target_invoice_id = Column(UUID(as_uuid=True), nullable=True) # Sales/Purchase Invoice
    target_cn_dn_id = Column(UUID(as_uuid=True), nullable=True) # Credit/Debit Note
    allocated_amount = Column(Float, nullable=False)
    is_floating = Column(Boolean, default=False)
    created_at = Column(DateTime, default=datetime.utcnow)""")
code.style.font.name = 'Courier New'

doc.add_heading('1.3 API Endpoints Required (FastAPI)', level=2)
doc.add_paragraph('- POST /api/finance/allocations : Creates a new allocation. MUST validate that allocated_amount does not exceed the remaining balance of the target_invoice or the source_voucher.')
doc.add_paragraph('- GET /api/finance/allocations/pending : Fetches all floating/unallocated vouchers and unpaid invoices for a specific Party.')

# Section 2: Challan-to-Invoice Data Pipeline
doc.add_heading('2. Challan-to-Invoice Backend Pipeline', level=1)
doc.add_paragraph('1. Database Requirement: Add `linked_challan_id` to the `Invoice` model.')
doc.add_paragraph('2. API Rule: When converting a Challan to an Invoice, the backend MUST NOT deduct physical inventory again (as the Challan already deducted it). It must only compute GST and hit the Financial Ledgers (Accounts Receivable & Sales).')

# Section 3: CRM Permissions Matrix
doc.add_heading('3. CRM & Multi-Tenant Permissions Matrix', level=1)
doc.add_paragraph('1. Database Requirement: Add a `role_permissions` JSONB column to the `UserRole` mapping table to handle granular module access.')
doc.add_paragraph('2. API Rule: Every protected endpoint MUST check the JWT token\'s `organization_id` (CM boundary) and the specific `role_permissions` JSON matrix before executing.')

doc.save('final implemantation of accounting and billing/Backend_Task_Specs_For_Claude.docx')
print('Successfully generated Backend_Task_Specs_For_Claude.docx')
