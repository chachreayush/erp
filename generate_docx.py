import sys
import subprocess

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# Title
title = doc.add_heading('Enterprise ERP Finance & Real-World Accounting Workflow', level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('A detailed blueprint of top-tier ERP financial architectures combined with practical, day-to-day accounting workflows for a modern firm.')

# Part 1
doc.add_heading('Part 1: The Best-in-Class ERP Financial System Workflow', level=1)
doc.add_paragraph('Top-tier ERPs (like SAP S/4HANA or Oracle NetSuite) operate on a unified data model. The workflow of data within the system is designed to eliminate manual data entry, reduce reconciliation, and automate reporting.')

doc.add_heading('A. The Universal Journal Data Flow', level=2)
p = doc.add_paragraph()
p.add_run('In the best ERPs, data flows into a single massive repository (The Universal Journal). When an operational event occurs (e.g., a warehouse worker scans a barcode to ship a product), the system automatically translates that physical action into a financial action. It simultaneously updates the inventory ledger, the customer ledger, and the general ledger without human intervention.')

doc.add_heading('B. Automated Subledger Syncing', level=2)
doc.add_paragraph("Subledgers (Accounts Payable, Accounts Receivable, Fixed Assets) are directly hard-wired to the General Ledger control accounts. A user cannot post a generic journal entry directly to 'Accounts Receivable'. They must process a valid customer invoice, ensuring the detailed customer balances always match the company's grand total.")

doc.add_heading('C. Continuous Consolidation & Multi-Currency', level=2)
doc.add_paragraph('Every financial transaction is recorded in three currencies simultaneously: the Transaction Currency (e.g., EUR), the Base Currency of the local branch (e.g., INR), and the Global Group Currency (e.g., USD). This workflow allows the parent company to instantly view real-time consolidated reports across all global branches without waiting for end-of-month manual conversions.')

# Part 2
doc.add_heading('Part 2: Real-World Firm Accounting Workflows', level=1)
doc.add_paragraph('Below is the step-by-step workflow of how a firm actually operates day-to-day using this modern ERP architecture. We will cover the two most critical business cycles: Procure-to-Pay (Buying things) and Order-to-Cash (Selling things).')

doc.add_heading('Cycle 1: Procure-to-Pay (P2P) - Purchasing & AP', level=2)
doc.add_paragraph('This is the workflow a firm uses to buy inventory or services and pay its vendors.')

doc.add_paragraph('Step 1: Purchase Requisition & Purchase Order (PO)', style='List Bullet')
doc.add_paragraph('Action: A department requests materials. Purchasing approves it and sends a PO to the vendor.')
doc.add_paragraph('Accounting Impact: None yet. This creates a "commitment" in the system so management knows money will be spent soon, but no financial ledgers are hit.')

doc.add_paragraph('Step 2: Goods Receipt (GR)', style='List Bullet')
doc.add_paragraph('Action: The warehouse receives the goods from the vendor and scans them into the ERP.')
doc.add_paragraph('Accounting Impact: The ERP automatically generates a financial entry behind the scenes.\n  - Debit (Increase): Inventory Asset Account\n  - Credit (Increase): GR/IR (Goods Received/Invoice Not Received) Liability Clearing Account')

doc.add_paragraph('Step 3: Vendor Invoice Receipt', style='List Bullet')
doc.add_paragraph('Action: The vendor emails the bill. The AP clerk enters it into the ERP. The ERP performs a "3-Way Match" (Checking if PO = Goods Received = Invoice).')
doc.add_paragraph('Accounting Impact:\n  - Debit (Decrease): GR/IR Clearing Account (Zeroing it out)\n  - Credit (Increase): Accounts Payable (Vendor Subledger)')

doc.add_paragraph('Step 4: Vendor Payment', style='List Bullet')
doc.add_paragraph('Action: Finance runs a weekly payment batch. The ERP sends funds via bank integration.')
doc.add_paragraph('Accounting Impact:\n  - Debit (Decrease): Accounts Payable\n  - Credit (Decrease): Cash/Bank Account')

doc.add_heading('Cycle 2: Order-to-Cash (O2C) - Sales & AR', level=2)
doc.add_paragraph('This is the workflow a firm uses to sell goods to customers and collect cash.')

doc.add_paragraph('Step 1: Sales Order', style='List Bullet')
doc.add_paragraph('Action: A customer agrees to buy products. A Sales Order is created.')
doc.add_paragraph('Accounting Impact: None yet. The ERP simply "allocates" the inventory so no one else buys it.')

doc.add_paragraph('Step 2: Fulfillment / Goods Issue', style='List Bullet')
doc.add_paragraph('Action: The warehouse packs and ships the box to the customer.')
doc.add_paragraph('Accounting Impact: The ERP instantly recognizes the cost of the sale.\n  - Debit (Increase): Cost of Goods Sold (COGS) Expense Account\n  - Credit (Decrease): Inventory Asset Account')

doc.add_paragraph('Step 3: Customer Billing / Invoice', style='List Bullet')
doc.add_paragraph('Action: The system automatically generates and emails an invoice to the customer based on the shipment.')
doc.add_paragraph('Accounting Impact: The ERP recognizes the revenue.\n  - Debit (Increase): Accounts Receivable (Customer Subledger)\n  - Credit (Increase): Sales Revenue Account\n  - Credit (Increase): Taxes Payable (GST/VAT)')

doc.add_paragraph('Step 4: Customer Payment Receipt', style='List Bullet')
doc.add_paragraph('Action: The customer pays the bill via bank transfer. The ERP bank feed auto-matches the deposit to the invoice.')
doc.add_paragraph('Accounting Impact:\n  - Debit (Increase): Cash/Bank Account\n  - Credit (Decrease): Accounts Receivable')

doc.add_heading('Cycle 3: Record-to-Report (R2R) - The Month-End Close', level=2)
doc.add_paragraph('Because the P2P and O2C cycles are fully automated and mathematically locked, the month-end close becomes a fast "soft close".')

doc.add_paragraph('1. Daily Bank Reconciliation:', style='List Number')
doc.add_paragraph('Instead of waiting until month-end, the ERP connects directly to the bank feed and auto-reconciles daily cash movements.')

doc.add_paragraph('2. Automated Depreciation:', style='List Number')
doc.add_paragraph('The ERP Fixed Asset module automatically calculates and posts the monthly depreciation expense for all equipment.')

doc.add_paragraph('3. One-Click Financial Statements:', style='List Number')
doc.add_paragraph('Because every transaction has been instantly hitting the Universal Journal all month, the CFO can click a button on the last day of the month to generate the Balance Sheet, Profit & Loss Statement, and Cash Flow Statement in real-time.')

doc.save(r'C:\Users\DELL\OneDrive\Desktop\erp2\ERP_Finance_Workflow_Guide.docx')
print('Document generated successfully.')
