import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

def create_doc10():
    doc = docx.Document()
    
    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("DOC-10\nCUSTOMER, SUPPLIER & PARTY MASTER ARCHITECTURE")
    run.font.name = 'Arial'
    run.font.size = Pt(16)
    run.font.bold = True
    
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s_run = subtitle.add_run("Implementation Specification • Phase 4: Master Data")
    s_run.font.name = 'Arial'
    s_run.font.size = Pt(12)
    s_run.font.italic = True
    
    doc.add_page_break()
    
    # TOC
    doc.add_heading("Document Map", level=1)
    toc_items = [
        "1. Executive Scope & Objectives",
        "2. Design Principles",
        "3. Master Data Philosophy",
        "4. Unified Party Architecture vs Separate Masters",
        "5. Customer Master Data Model",
        "6. Supplier Master Data Model",
        "7. Transporter & Logistics Party Model",
        "8. Multi-Branch & Multi-Address Handling",
        "9. Statutory & Compliance Details (GST, PAN, FSSAI)",
        "10. Financial Linkage (General Ledger AR/AP)",
        "11. Credit Control & Limits Engine",
        "12. Pricing, Discount & Scheme Defaults",
        "13. Route, Territory & Field Staff Mapping",
        "14. Master Document Attachments (KYC)",
        "15. Active/Inactive & Archiving Lifecycle",
        "16. Duplicate Prevention & Resolution",
        "17. Data Privacy & Tenant Isolation",
        "18. Database Entities & Relationships",
        "19. API / Service Contracts",
        "20. Keyboard-First UX Workflows",
        "21. Small vs Large Business Modes",
        "22. Permissions, Audit Trail & Security",
        "23. Local-First Synchronization Rules",
        "24. Failure & Edge Cases",
        "25. Implementation Checklist"
    ]
    for item in toc_items:
        doc.add_paragraph(item)
        
    doc.add_page_break()

    sections = {
        "1. Executive Scope & Objectives": "DOC-10 defines the comprehensive architectural specification for the Customer, Supplier, and Transporter masters (collectively 'Party Master'). It establishes the core foundational entities required before sales billing, purchasing, or open-item accounting can occur. This document integrates strictly with the accounting rules set forth in DOC-09 (AR/AP Open Items) and the CRM/Field Staff architecture.",
        
        "2. Design Principles": "1. Single Source of Truth: A party exists once globally within a tenant; modules reference it rather than duplicating it.\n2. Role-Based Expansion: A single party can act as both a Customer and a Supplier without duplicating their fundamental KYC data.\n3. Zero Duplicate Entry: GST numbers and PAN numbers are strictly validated globally to prevent dirty data.\n4. Strong Financial Ties: Every party must map to a specific General Ledger Subledger Account.\n5. Immutable Financial Identity: Even if a customer is deactivated, their financial history remains perfectly intact and auditable.",
        
        "3. Master Data Philosophy": "The ERP must treat Master Data as the supreme authority. Transactions are highly volatile, but Masters are highly stable. Creation of Master Data requires elevated permissions (CM Admin / Master Admin) and goes through an audit trail. All default configurations (pricing, tax calculations, credit limits) flow downhill from the Master to the Transaction.",
        
        "4. Unified Party Architecture vs Separate Masters": "The system will utilize a Base Party architectural pattern. A 'Party' holds the core identity (Legal Name, PAN, GST). The 'Customer Profile' and 'Supplier Profile' are extensions of this base identity. This prevents the common ERP flaw where 'ABC Corp' exists as both Customer #105 and Supplier #402 with desynced addresses.",
        
        "5. Customer Master Data Model": "Customer extensions include: Sales Territory mapping, Default Price List, Assigned Field Sales Representative, Default Transporter, Route scheduling data, Delivery windows, and Customer Class/Segment (Retailer, Distributor, Wholesaler).",
        
        "6. Supplier Master Data Model": "Supplier extensions include: Purchasing terms, Default Manufacturer/Principal mapping, Minimum Order Quantities (MOQ), Lead Time, Supplier Rating, and Default Purchase Ledger Account.",
        
        "7. Transporter & Logistics Party Model": "Since this is a C&F-native ERP, Transporters are treated as first-class masters, not just text fields. Extensions include: Fleet tracking details, Default Freight Rates (Per KG, Per Box), Hub locations, and logistics recovery accounts.",
        
        "8. Multi-Branch & Multi-Address Handling": "A single party can have multiple physical locations. The architecture supports a 1-to-N relationship between Party and Address. Addresses are strictly categorized as: Billing Address, Shipping Address, or Corporate Address. Sales Orders must dynamically select the Shipping Address without altering the Billing Address tied to the GST.",
        
        "9. Statutory & Compliance Details (GST, PAN, FSSAI)": "Strict validation logic must be enforced on input. GST lengths must be 15 characters, PAN must be 10 alphanumeric. The system must support Drug License Numbers and FSSAI numbers with automatic expiration alerts for medical C&F operations.",
        
        "10. Financial Linkage (General Ledger AR/AP)": "Inheriting rules from DOC-09, every new Customer is automatically mapped to an AR Subledger under 'Sundry Debtors'. Every new Supplier maps to an AP Subledger under 'Sundry Creditors'. The user must never manually type ledger codes; the system handles the double-entry binding.",
        
        "11. Credit Control & Limits Engine": "The system will strictly enforce two limits: Credit Amount (Max ₹ outstanding) and Credit Days (Max days outstanding). During the Sales Billing process (DOC-22), the system will actively block or flag transactions that exceed these limits based on the Permission Engine (DOC-04).",
        
        "12. Pricing, Discount & Scheme Defaults": "Customers can be tagged with Default Discount Percentages or assigned to specific Pricing Tiers (e.g., Tier A, Tier B). This eliminates manual pricing errors during high-speed billing operations.",
        
        "13. Route, Territory & Field Staff Mapping": "Customers are organized geographically. The system maps Customers to Territories, Territories to Routes, and Routes to specific Field Staff. This enables the Field Staff Tracking & Scheduling module.",
        
        "14. Master Document Attachments (KYC)": "Users can upload KYC PDFs/Images (GST Certificates, Drug Licenses) directly to the Master Profile. Files are stored securely using the Cloud Document Storage architecture, linked by Party ID.",
        
        "15. Active/Inactive & Archiving Lifecycle": "Parties cannot be deleted once they have associated financial transactions. They can only be marked as 'Inactive'. Inactive parties are hidden from Dropdowns and Lookups to maintain UI speed but remain visible in historical reports.",
        
        "16. Duplicate Prevention & Resolution": "The UI will perform an asynchronous check while typing the GST/PAN or Mobile Number. If a match is found, the UI blocks creation and prompts the user to merge or edit the existing party.",
        
        "17. Data Privacy & Tenant Isolation": "As per DOC-01, every Party record contains an `organization_id`. API endpoints will strictly filter by this ID to guarantee absolute tenant isolation.",
        
        "18. Database Entities & Relationships": "Table: `parties` (id, org_id, legal_name, pan, gst, status)\nTable: `party_addresses` (id, party_id, type, line1, city, state, pincode)\nTable: `customers` (party_id, route_id, credit_limit, price_list)\nTable: `suppliers` (party_id, payment_terms, lead_time)",
        
        "19. API / Service Contracts": "Standard RESTful endpoints:\n`POST /api/master/parties`\n`GET /api/master/customers`\n`PUT /api/master/parties/{id}`\nPayloads must include nested addresses and be validated via Pydantic schemas.",
        
        "20. Keyboard-First UX Workflows": "Following DOC-05, the Master Creation screen allows the user to press 'Tab' rapidly through fields, hit 'Enter' to save, and immediately return to a blank form for high-speed data entry. 'Ctrl+F' globally searches parties.",
        
        "21. Small vs Large Business Modes": "Small businesses see a single flat form. Large businesses see tabs (General, Addresses, Financials, CRM, Statutory) to handle complex, multi-departmental data.",
        
        "22. Permissions, Audit Trail & Security": "Changes to Credit Limits or Statutory fields log an explicit Audit Event (Old Value → New Value, User, Timestamp) to prevent fraud.",
        
        "23. Local-First Synchronization Rules": "Master data changes made on a LAN node are synchronized to the Cloud Light Database asynchronously. Due to their low-velocity nature compared to transactions, Master Sync uses a standard Last-Write-Wins (LWW) conflict resolution strategy.",
        
        "24. Failure & Edge Cases": "If the GST API validation goes offline, the system allows 'Draft' creation of the Customer, enabling business to continue, but flags it for later verification.",
        
        "25. Implementation Checklist": "[ ] Database schema migrations\n[ ] Unified API routes\n[ ] Keyboard-first React UI\n[ ] Credit limit middleware\n[ ] Role-based permission integration"
    }

    for heading, text in sections.items():
        doc.add_heading(heading, level=1)
        doc.add_paragraph(text)

    # Save
    out_path = r'C:\Users\DELL\Desktop\New folder (3)\DOC-10_Customer_and_Supplier_Master_v1_0_RELEASE.docx'
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == "__main__":
    create_doc10()
