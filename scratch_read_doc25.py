import docx

doc = docx.Document(r"C:\Users\DELL\Desktop\New folder (3)\DOC-25_Journal_Voucher_Accounting_Adjustment_Accrual_Management_Engine_v1_0_FINAL.docx")
with open("doc25_summary.txt", "w", encoding="utf-8") as f:
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            f.write(f"{i}: {p.text}\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- TABLE {t_idx} ---\n")
        for row in table.rows:
            row_txt = [c.text.strip().replace('\n', ' ') for c in row.cells]
            f.write(" | ".join(row_txt) + "\n")

print("Saved doc25_summary.txt")
