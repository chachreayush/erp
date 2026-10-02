import os
from docx import Document
from docx.shared import Pt, RGBColor

FOLDER_PATH = r"C:\Users\DELL\Desktop\New folder (3)"

def add_appendix(doc, title, content_blocks):
    doc.add_page_break()
    heading = doc.add_heading(f"APPENDIX: ACTUAL IMPLEMENTED SOFTWARE FEATURES (AS OF LATEST BUILD)", level=1)
    
    doc.add_paragraph(
        "The following specifications represent the exact features, UI components, and database models "
        "that have been actively coded and deployed in the current erp2 software build. This section "
        "reconciles the theoretical architecture with the live working software, ensuring no details are lost."
    )
    
    for block_title, text in content_blocks.items():
        doc.add_heading(block_title, level=2)
        doc.add_paragraph(text)

def update_doc01():
    path = os.path.join(FOLDER_PATH, "DOC-01_Platform_Architecture_Company_Isolation_v1_0.docx")
    if not os.path.exists(path): return
    doc = Document(path)
    content = {
        "1. Multi-Tenant SaaS Engine (Admin & Client Management)": 
            "The software currently features a hard-isolated SaaS layer. The root organization is designated as 'AM-0001' (System Admin). "
            "Users belonging to AM-0001 have access to a dedicated 'Client Management' dashboard (/clients route).",
        "2. Client Creation UI": 
            "The UI features a full 'Client Management' interface where the Admin can click '+ Add Client' to provision new tenant companies. "
            "Fields include: Company Name, Subdomain/Org ID, Admin Email, Initial Password, and Database Schema routing. "
            "The system automatically provisions a unique UUID for the organization and seeds a default 'admin' user for that tenant.",
        "3. Database Isolation Mechanism": 
            "In models.py, every single operational table (User, Ledger, Voucher, Product, Party) requires an 'organization_id' Foreign Key. "
            "All API queries explicitly filter by 'current_user.organization_id' ensuring strict data boundaries."
    }
    add_appendix(doc, "DOC-01", content)
    doc.save(path)
    print("Updated DOC-01")

def update_doc10():
    path = os.path.join(FOLDER_PATH, "DOC-10_Customer_Supplier_Business_Partner_Master_Framework_v1_0_RELEASE.docx")
    if not os.path.exists(path): return
    doc = Document(path)
    content = {
        "1. Unified Party Modal UI": 
            "The UI operates at /master/parties. It features a modern dark-mode modal for creating Unified Business Partners. "
            "It captures Core Details (Legal Name, Trade Name, GSTIN, PAN).",
        "2. Dual-Role Accounting (AR/AP Auto-Link)": 
            "The system handles companies that are BOTH buyers and sellers. When 'Is Customer' and 'Is Supplier' checkboxes are ticked, "
            "the UI dynamically exposes TWO separate ledger linking sections. One maps the party to an Accounts Receivable (AR) Ledger Group, "
            "and the other to an Accounts Payable (AP) Ledger Group.",
        "3. Database Architecture": 
            "Models implemented: 'parties' (base entity), 'party_addresses' (1-to-many), 'customer_profiles' (holds credit limits and AR ledger_id), "
            "and 'supplier_profiles' (holds payment terms and AP ledger_id). The backend API automatically generates the underlying financial ledgers."
    }
    add_appendix(doc, "DOC-10", content)
    doc.save(path)
    print("Updated DOC-10")

def update_doc07():
    path = os.path.join(FOLDER_PATH, "DOC-07_Complete_Accounting_Bill_by_Bill_Settlement_Engine.docx")
    if not os.path.exists(path): return
    doc = Document(path)
    content = {
        "1. Financial Core (Finance V2)": 
            "The live software implements a robust Double-Entry accounting engine. The VoucherEntry UI supports dynamic multi-row debit/credit grids. "
            "It enforces strict balancing (Total Dr = Total Cr) before submission.",
        "2. Live Financial Reports": 
            "The backend dynamically calculates and the frontend renders: Day Book, Ledger Statements (with running balances), "
            "Trial Balance, Profit & Loss Statement, and a strictly structured Balance Sheet.",
        "3. Account Heads (Ledger Group Master)": 
            "The system implements standard Indian Taxation Account Heads (Assets, Liabilities, Income, Expenses). "
            "A visual hierarchical tree UI is implemented in LedgerGroupMaster.tsx showing nested groups (e.g. Current Assets -> Sundry Debtors)."
    }
    add_appendix(doc, "DOC-07", content)
    doc.save(path)
    print("Updated DOC-07")

def update_doc06():
    path = os.path.join(FOLDER_PATH, "DOC-06_Master_Data_and_Batch_Framework_v1_0_FINAL2.docx")
    if not os.path.exists(path): return
    doc = Document(path)
    content = {
        "1. Pharmaceutical Master Data": 
            "The MasterPage.tsx implements dedicated tabs and API-driven CRUD modals for: Salt Master (chemical compositions), "
            "Company Master (manufacturers), HSN Code Master (tax rates), and State Code Master (GST compliance).",
        "2. Keyboard-First Operations": 
            "The UI allows users to navigate the data grids using Arrow Keys, hit Enter to edit, and hit Escape to return to the dashboard, "
            "strictly adhering to high-speed billing requirements."
    }
    add_appendix(doc, "DOC-06", content)
    doc.save(path)
    print("Updated DOC-06")

try:
    update_doc01()
    update_doc10()
    update_doc07()
    update_doc06()
    print("All documents updated successfully.")
except Exception as e:
    print(f"Error: {e}")
