import docx

doc = docx.Document()
doc.add_heading('FRONTEND UI SPECIFICATION: Receipt Allocation Modal (For Claude)', 0)

doc.add_heading('1. Overview', level=1)
doc.add_paragraph('This document outlines the UI/UX requirements for the Bill-by-Bill Allocation Modal. This modal appears when a user saves a Receipt or Payment Voucher, allowing them to map the received money to outstanding invoices.')

doc.add_heading('2. Component: ReceiptAllocationModal.tsx', level=1)
doc.add_paragraph('Path: src/pages/finance/ReceiptAllocationModal.tsx')
doc.add_paragraph('Design: Use the dark enterprise theme consistent with the rest of the app (bg-slate-900, text-gray-100).')

doc.add_heading('2.1 UI Layout & Features', level=2)
doc.add_paragraph('1. Header: Display the Voucher Number and the Total Amount Received (e.g., "Allocating Receipt #RCT-001 (₹ 50,000)").')
doc.add_paragraph('2. Data Fetching: On mount, call GET /api/finance/allocations/pending to fetch outstanding invoices for the selected party.')
doc.add_paragraph('3. Grid Display: Show a table of outstanding invoices with columns: Date, Invoice No, Total Amount, Outstanding, and an "Allocate" input field.')
doc.add_paragraph('4. Floating Balance: At the bottom, dynamically calculate and display: "Total Received" - "Total Allocated" = "Floating / On Account Balance".')
doc.add_paragraph('5. Submission: On click of "Save Allocations", construct the AllocationCreate payload and POST it to /api/finance/allocations. If successful, close the modal.')

doc.add_heading('3. Integration with VoucherEntry.tsx', level=1)
doc.add_paragraph('Update the existing Voucher entry screen. When a Receipt or Payment is successfully saved, instead of just returning to the dashboard, pop open the `ReceiptAllocationModal`.')

doc.save('final implemantation of accounting and billing/Frontend_Allocation_Spec_For_Claude.docx')
print('Successfully generated Frontend_Allocation_Spec_For_Claude.docx')
