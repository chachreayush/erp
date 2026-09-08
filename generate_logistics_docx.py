import sys
import os

try:
    from docx import Document
    from docx.shared import Pt, RGBColor
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, '-m', 'pip', 'install', 'python-docx'])
    from docx import Document
    from docx.shared import Pt, RGBColor

folder_path = r'C:\Users\DELL\OneDrive\Desktop\erp2\finance and accounting document'
file_path = os.path.join(folder_path, 'ERP_Logistics_and_Pricing_Architecture.docx')

doc = Document()
doc.add_heading('Enterprise Logistics, Pricing & Financial Integration Architecture', level=0)
doc.add_paragraph('A deep-dive companion guide detailing how Commercial Intent, Physical Logistics, and Financial Accounting are architecturally decoupled and integrated, based on SAP S/4HANA (SD/MM) and Oracle NetSuite.')

# 1. Separation of Concerns
doc.add_heading('1. The Separation of Concerns', level=1)
doc.add_paragraph('Top-tier ERPs strictly separate the lifecycle of a transaction into three distinct layers:')
doc.add_paragraph('Commercial Intent (Sales Orders / POs): Non-posting agreements. They allocate stock and budget but hit zero financial ledgers.', style='List Bullet')
doc.add_paragraph('Physical Logistics (Deliveries / Goods Receipts): Alters physical inventory and triggers Inventory Valuation journals (e.g., hitting Cost of Goods Sold).', style='List Bullet')
doc.add_paragraph('Financial Accounting (Billing / Invoices): Establishes the legal liability, hitting Accounts Receivable/Payable, Revenue, and Tax.', style='List Bullet')

# 2. O2C Flow & ATP
doc.add_heading('2. Order-to-Cash (O2C) & Available-to-Promise (ATP)', level=1)
doc.add_paragraph('When a Sales Order is created, an ATP engine dynamically allocates stock without creating static silos. The critical financial trigger is the Post Goods Issue (PGI) when the box leaves the warehouse. PGI triggers a Debit to COGS and a Credit to Inventory. Note: PGI recognizes Cost of Sales, NEVER Revenue. Revenue is deferred until the final Billing Invoice is generated.')

# 3. P2P Flow & GR/IR Clearing
doc.add_heading('3. Procure-to-Pay (P2P) & GR/IR Clearing', level=1)
doc.add_paragraph('When goods arrive before the vendor invoice, ERPs use a Suspense Architecture via a GR/IR (Goods Receipt / Invoice Receipt) Clearing Account.')
doc.add_paragraph('At Goods Receipt (Dock): Debit Inventory Asset, Credit GR/IR Clearing (Provisional Liability).', style='List Number')
doc.add_paragraph('At Invoice Receipt (Desk): A 3-Way Match occurs. The system Debits GR/IR Clearing (zeroing it out), Debits Input Tax, and Credits Accounts Payable.', style='List Number')

# 4. Document Lineage & DAG
doc.add_heading('4. Document Lineage & DAG Architecture', level=1)
doc.add_paragraph('Enterprise systems treat transactions as a Directed Acyclic Graph (DAG). Documents are never deleted. If a mistake is made, the system enforces "Storno Accounting" (Negative Postings) or Compensating Workflows (Return Orders, Credit Memos) to reverse the flow, guaranteeing a 100% immutable audit trail.')

# 5. Pricing & Condition Technique
doc.add_heading('5. The Condition Technique (Pricing & Tax Engine)', level=1)
doc.add_paragraph('Pricing is decoupled from product master data using an advanced pipeline:')
doc.add_paragraph('Access Sequence: The system searches from specific (Customer + Product rate) to generic (Base Product rate) until it finds a hit.', style='List Bullet')
doc.add_paragraph('Account Determination Bridge: Instead of hardcoding GL accounts to products, the pricing engine tags lines with "Account Keys" (e.g., ERL for Revenue, MWS for Tax). A matrix then dynamically determines the correct GL account at the moment of posting, allowing sales teams to change pricing rules without breaking the General Ledger.', style='List Bullet')

doc.save(file_path)
print('Logistics & Pricing Companion Document generated successfully!')
