import sys
import subprocess
try:
    from docx import Document
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document

with open('reports.txt', 'r', encoding='utf-8') as f:
    text = f.read()

doc = Document()
doc.add_heading('Reports & GST Module Implementation', 0)

for line in text.split('\n'):
    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('# '):
        doc.add_heading(line[2:], level=1)
    else:
        doc.add_paragraph(line)

doc.save('reports.docx')
