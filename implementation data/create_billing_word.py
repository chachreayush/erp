import sys
import subprocess
try:
    from docx import Document
    from docx.shared import Inches
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    from docx import Document
    from docx.shared import Inches

with open('billing.txt', 'r', encoding='utf-8') as f:
    text = f.read()

doc = Document()
doc.add_heading('Billing & Consignment Architecture', 0)

# Add the UI Image at the very top of the document
try:
    image_path = r'C:\Users\DELL\.gemini\antigravity\brain\8af5d42e-53e3-4425-a748-cbfe8c6eb22c\.user_uploaded\media_1789338431515.png'
    doc.add_picture(image_path, width=Inches(6.0))
    doc.add_paragraph('UI Reference: Sales Invoice Dashboard Layout')
except Exception as e:
    doc.add_paragraph(f'[Error loading image: {str(e)}]')

# Parse the text file and add headings correctly
for line in text.split('\n'):
    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
    elif line.startswith('## '):
        doc.add_heading(line[3:], level=2)
    elif line.startswith('# '):
        doc.add_heading(line[2:], level=1)
    else:
        doc.add_paragraph(line)

doc.save('billing.docx')
