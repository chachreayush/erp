import docx
import os

filepath = 'c:/Users/DELL/OneDrive/Desktop/erp2/AI_Collaboration_Protocol.docx'

if os.path.exists(filepath):
    doc = docx.Document(filepath)
else:
    doc = docx.Document()
    doc.add_heading('AI COLLABORATION PROTOCOL & CREDIT CONSERVATION RULES', 0)

# Section 3: Preventing Exhaustion & Incomplete Work
doc.add_heading('3. Preventing Claude Credit Exhaustion (Chunking Strategy)', level=1)
doc.add_paragraph('To ensure that our work is never left incomplete if Claude suddenly runs out of credits, Gemini will enforce the following strategy:')
doc.add_paragraph('- Micro-Tasking: Gemini will break down every major feature into tiny, independent coding tasks.')
doc.add_paragraph('- One Prompt, One Feature: Gemini will only tell the user to switch to Claude when a task is small enough to be completed in exactly 1 or 2 prompts.')
doc.add_paragraph('- Checkpointing: After every small chunk of code is returned by Claude, Gemini will immediately test it and commit it to git. If Claude exhausts its credits, the project will still be in a stable, runnable state rather than a half-broken state.')

doc.save(filepath)
print('Successfully updated AI_Collaboration_Protocol.docx')
