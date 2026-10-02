import docx
from docx.shared import Pt, RGBColor

doc = docx.Document()

# Title
title = doc.add_heading('AI COLLABORATION PROTOCOL & CREDIT CONSERVATION RULES', 0)

# Section 1: The AI Roles
doc.add_heading('1. The Division of Labor', level=1)
doc.add_paragraph('To maximize efficiency and conserve limited AI credits, the development of this ERP project follows a strict division of labor between Gemini (Antigravity) and Claude.')

doc.add_heading('1.1 Claude\'s Role (The Coder)', level=2)
doc.add_paragraph('Claude is strictly assigned to executing the actual CODE. Whether it is Backend (FastAPI, Python, SQLAlchemy) or Frontend (React, Vite, UI/UX), Claude writes the code based on the blueprints provided.')
doc.add_paragraph('By restricting Claude to only coding, we drastically conserve its limited token/credit quota.')

doc.add_heading('1.2 Gemini\'s Role (The Architect & Tester)', level=2)
doc.add_paragraph('Gemini handles all heavy lifting that does not involve writing production code. This includes:')
doc.add_paragraph('- Testing the application (running tests, reading error logs).')
doc.add_paragraph('- Providing detailed feedback on code written by Claude.')
doc.add_paragraph('- Creating technical outlines, blueprints, and .docx implementation specs.')
doc.add_paragraph('- Running automated visual audits via Selenium scripts.')
doc.add_paragraph('- Managing the project workflow and directory synchronization.')

# Section 2: The Workflow Handoff
doc.add_heading('2. When to Switch AI Models', level=1)
doc.add_paragraph('Gemini acts as the Project Manager. Gemini will explicitly tell the user when it is time to switch to Claude, and when to switch back. The flow is as follows:')
doc.add_paragraph('Step 1 [Gemini]: Gemini creates the detailed .docx blueprint/outline for a feature.')
doc.add_paragraph('Step 2 [Claude]: The user gives the blueprint to Claude. Claude writes the code and implements the feature.')
doc.add_paragraph('Step 3 [Gemini]: The user switches back to Gemini. Gemini tests the feature, runs the servers, performs visual audits, and provides feedback/fixes.')
doc.add_paragraph('Step 4 [Claude]: Claude implements any necessary fixes based on Gemini\'s feedback.')

doc.save('c:/Users/DELL/OneDrive/Desktop/erp2/AI_Collaboration_Protocol.docx')
print('Successfully generated AI_Collaboration_Protocol.docx')
