import os
from docx import Document

FOLDER_PATH = r"C:\Users\DELL\Desktop\New folder (3)"

def append_to_doc(filename, title, text):
    path = os.path.join(FOLDER_PATH, filename)
    if not os.path.exists(path): return
    doc = Document(path)
    
    has_tier = any("Multi-Tiered Admin" in p.text for p in doc.paragraphs)
    if not has_tier:
        doc.add_heading(title, level=2)
        doc.add_paragraph(text)
        doc.save(path)
        print(f"Updated {filename}")
    else:
        print(f"Skipped {filename}, already updated.")

doc01 = "DOC-01_Platform_Architecture_Company_Isolation_v1_0.docx"
text01 = (
    "In addition to standard Client Companies, the architecture supports the creation of "
    "additional Admin Companies (e.g., Resellers or Regional Master Admins). A Root System Admin "
    "can provision a new 'Admin Company', which inherits the permission to create and manage its own "
    "sub-tier of Client Companies. Furthermore, multiple Admin Users can be created within any Admin Company "
    "to distribute platform management workloads. (Note: While the database structure supports this hierarchy, "
    "the frontend UI for creating new Admin Companies is scheduled for a future release)."
)

doc03 = "DOC-03_Users_Roles_Permissions_Engine_v1_0_RELEASE.docx"
text03 = (
    "The role engine explicitly supports a Multi-Tiered Admin Hierarchy. At the highest level is the "
    "Root System Admin. This admin can create both 'Client Companies' and new 'Admin Companies'. "
    "Users assigned to an Admin Company possess platform-level privileges to provision and oversee clients "
    "assigned to them, effectively enabling a white-label or reseller ecosystem. The engine also allows adding "
    "multiple individual 'Admin Users' to share the workload."
)

append_to_doc(doc01, "4. Multi-Tiered Admin & Reseller Architecture", text01)
append_to_doc(doc03, "APPENDIX: Admin Company & Reseller Capabilities", text03)
